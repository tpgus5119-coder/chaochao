import re, glob, json, zipfile, unicodedata, collections, sys, time
import openpyxl
nfc=lambda s: unicodedata.normalize('NFC', str(s)).strip().lower()
VI=re.compile(r'[ăâđêôơưàảãáạằẳẵắặầẩẫấậèẻẽéẹềểễếệìỉĩíịòỏõóọồổỗốộờởỡớợùủũúụừửữứựỳỷỹýỵ]')
found={}
def put(w,kind,no):
    k=nfc(w)
    if len(k)<2: return
    if k not in found: found[k]=(kind,no)
files=sorted(glob.glob('*'))
for f in files:
    t0=time.time()
    m=re.search(r'(\d+)\s*\)?\s*회차', f); w=re.search(r'(\d+)\s*주차', f)
    kind,no=('일일',int(m.group(1))) if m else ('주간',int(w.group(1))) if w else (None,None)
    if not kind: print('skip', f, flush=True); continue
    try:
        if f.endswith('.xlsx'):
            wb=openpyxl.load_workbook(f, read_only=True, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 3000), max_col=min(ws.max_column or 0, 40), values_only=True):   # 서식 때문에 94,238×16,383 으로 잡힌 시트가 있다 — 단어 목록은 앞쪽 몇천 줄뿐
                    for c in row:
                        if isinstance(c,str) and VI.search(c) and len(c)<40: put(c,kind,no)
            wb.close()
        elif f.endswith('.docx'):
            z=zipfile.ZipFile(f); x=z.read('word/document.xml').decode('utf8'); x=re.sub(r'</w:p>','\n',x); x=re.sub(r'<[^>]+>','',x)
            for line in x.split('\n'):
                for piece in re.split(r'[\t/;,()]|\s{2,}', line):
                    piece=piece.strip()
                    if VI.search(piece) and 1<len(piece)<40 and not re.search(r'[가-힣A-Za-z]{3,}\s*$', piece) and not re.search(r'[가-힣]', piece): put(piece,kind,no)
        else: print('skip', f, flush=True); continue
    except Exception as e: print('err', f, repr(e)[:60], flush=True); continue
    print(f'{f} {kind}{no} {time.time()-t0:.1f}s 누적 {len(found)}', flush=True)
json.dump({w:list(v) for w,v in found.items()}, open('/Users/leesehyeon/짜오짜오/베트남어-어플/scratchpad/c20_word_rounds.json','w',encoding='utf-8'), ensure_ascii=False)
S=json.load(open('/Users/leesehyeon/짜오짜오/베트남어-어플/data/_senior_words.json'))
kit=[s for s in S['sets'] if s['kind']=='기타'][0]
hit=[(w['vi'],found[nfc(w['vi'])]) for w in kit['words'] if nfc(w['vi']) in found]
print('기타', len(kit['words']), '중 회차를 찾은 것', len(hit), collections.Counter(k for _,(k,_) in hit), flush=True)
print('끝', flush=True)
