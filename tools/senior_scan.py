#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""기수 폴더의 **모든 파일**을 열어 무엇이 들었는지 적는다 → data/_senior_scan-{기수}.json

왜 만들었나 (2026-08-30): 이름만 보고 거른 탓에 회차를 놓쳤다.
  · 「답안」이라 적혀 있어도 **뜻 칸이 비어 있는** 판이 있다(17기). 낱말만 건진다.
  · 뜻이 **영어**로 된 판이 있다(18기 8·9회차 등). 버리지 말고 영어를 받아 둔다.
  · 스캔 그림은 글자층이 없다 — 이것만 진짜로 못 쓴다.
쓰기: python3 tools/senior_scan.py 17 18 19 20
"""
import json, os, pathlib, re, sys, collections
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import senior18 as S

R = pathlib.Path(__file__).resolve().parent.parent
EN = re.compile(r"^[A-Za-z][A-Za-z .,'\-/()]*$")

def rows_xlsx_lean(p):
    """15MB 짜리 엑셀이 있다 — 통째로 열면 메모리가 터진다. 읽기전용 + 줄 제한."""
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    for ws in wb.worksheets:
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > 4000: break
            cells = [("" if c is None else str(c).strip()) for c in r[:12]]
            if any(cells): out.append(cells)
    wb.close()
    return out


def cells_of(p):
    f = p.name.lower()
    if f.endswith((".xlsx", ".xls")): return rows_xlsx_lean(p)
    if f.endswith(".docx"):
        import docx
        out = []
        d = docx.Document(str(p))
        for t in d.tables:
            for row in t.rows:
                out.append([c.text.strip() for c in row.cells])
        for para in d.paragraphs:
            if para.text.strip(): out.append([para.text.strip()])
        return out
    return S.rows_pdf(p)

def split(rows):
    """줄마다 (베트남어, 한글뜻, 영어뜻) 을 낸다. 없으면 빈 칸."""
    out = []
    for line in rows:
        cells = [c.strip() for c in line if c and c.strip()]
        if len(cells) < 2: continue
        j = " ".join(cells)
        if re.match(r"^\s*(No\.?|단어|뜻|날짜|이름|번호|STT)", j): continue
        ko = [c for c in cells if S.KO.search(c)]
        rest = [c for c in cells if not S.KO.search(c) and re.search(r"[A-Za-zÀ-ỹ]", c)
                and not re.fullmatch(r"[\d.\s]+", c)]
        vi = [c for c in rest if S.VI.search(c)]
        en = [c for c in rest if not S.VI.search(c) and EN.match(c)]
        if not vi and rest: vi = rest[:1]; en = [c for c in rest[1:] if EN.match(c)]
        if not vi: continue
        out.append((" ".join(vi).strip(), " / ".join(ko).strip(), " / ".join(en).strip()))
    return out

def main():
    for gi in sys.argv[1:]:
        D = [x for x in S.BASE.iterdir() if x.is_dir() and x.name.startswith(gi)][0]
        recs, why = [], collections.Counter()
        for f in sorted(os.listdir(D)):
            if f.startswith(".") or not f.lower().endswith((".pdf",".xlsx",".xls",".docx")):
                why["시험지 아님(소리·그림 등)"] += 1; continue
            try: rows = cells_of(D / f)
            except Exception as e: why[f"못 읽음 {type(e).__name__}"] += 1; continue
            if not rows: why["스캔 그림(글자층 없음)"] += 1; continue
            ws = split(rows)
            if len(ws) < 5: why["낱말 5개 미만"] += 1; continue
            nk = sum(1 for _, k, _ in ws if k)
            ne = sum(1 for _, _, e in ws if e)
            tone = sum(1 for v, _, _ in ws if S.VI.search(v)) / len(ws)
            kind = ("한글" if nk >= len(ws) * .5 else "영어" if ne >= len(ws) * .5 else "낱말만")
            if tone < .25 and len(ws) >= 15: why["베트남어 아님"] += 1; continue
            lab = S.label(f) or ("일일", 0)
            recs.append({"src": f, "kind": lab[0], "no": lab[1], "meaning": kind,
                         "rows": [{"vi": v, "ko": k, "en": e} for v, k, e in ws]})
        (R / "data" / f"_senior_scan-{gi}.json").write_text(
            json.dumps({"gi": gi, "files": recs}, ensure_ascii=False, indent=1), encoding="utf-8")
        c = collections.Counter(r["meaning"] for r in recs)
        print(f"{gi}기 · 읽은 파일 {len(recs)} · 낱말자리 {sum(len(r['rows']) for r in recs)} · 뜻갈래 {dict(c)}")
        print("   못 쓴 파일:", dict(why))

main()
